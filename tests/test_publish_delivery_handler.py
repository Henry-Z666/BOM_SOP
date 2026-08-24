from __future__ import annotations

from pathlib import Path
from hashlib import sha256
import tempfile
import unittest

from PIL import Image
from openpyxl import load_workbook

from sop_pipeline.agent.artifacts import ArtifactStore
from sop_pipeline.agent.models import (
    RunRecord,
    RunStatus,
    StepResult,
    StepStatus,
)
from sop_pipeline.agent.review import SKIP_PUBLICATION_CANDIDATE_ID
from sop_pipeline.agent.skill_handlers import publish_delivery
from sop_pipeline.agent.skill_registry import SkillInvocation
from sop_pipeline.agent.skill_runtime import SkillContext
from sop_pipeline.agent.store import RunStore


def _formal_step(step_id: str, row: int) -> dict:
    return {
        "step_id": step_id,
        "main_process_id": "process-001",
        "title": step_id,
        "source_bom_rows": [row],
        "stage_scope_occurrence": "ROOT",
        "moving_occurrences": [],
        "receiver_occurrences": [],
        "visible_occurrences": [],
        "constraint_ids": [],
        "receiver_point_root": None,
        "receiver_normal_root": None,
        "translation_vector_root": None,
        "arrow_anchors": [],
        "camera_id": "fixed_123",
        "allowed_camera_ids": ["fixed_123", "fixed_456"],
        "depends_on": [],
        "affected_descendants": [],
        "state_delta": [],
        "complete_state_hash": f"sha256:state-{step_id}",
        "status": "ready",
        "diagnostics": [],
    }


def _bom_row(row: int, code: str, name: str) -> dict:
    return {
        "row": row,
        "level": "30" if row == 1 else f"30.{row - 1}",
        "material_code": code,
        "drawing_no": code,
        "name": name,
        "model": "",
        "quantity": 1,
        "unit": "件",
        "assembly_text": f"安装{name}",
        "control_points": "确认到位",
        "tools": "常规工具",
        "main_process_number": 1,
        "process_only": row == 1,
    }


class PublishDeliveryHandlerTests(unittest.TestCase):
    def test_later_resolution_preserves_previously_accepted_step(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workspace = root / "runs" / "run-1"
            rendered = workspace / "rendered"
            rendered.mkdir(parents=True)
            for step_id, shade in (
                ("step-1", 220),
                ("step-2", 200),
                ("step-3", 180),
                ("step-3-revision-0001", 120),
            ):
                Image.new("RGB", (320, 240), (shade, shade, shade)).save(
                    rendered / f"{step_id}.jpg"
                )
            revised_step_3_hash = "sha256:" + sha256(
                (rendered / "step-3-revision-0001.jpg").read_bytes()
            ).hexdigest()

            store = RunStore(root / "agent.sqlite3")
            bom_file = root / "BOM.xlsx"
            bom_file.write_bytes(b"bom")
            cad_directory = root / "cad"
            cad_directory.mkdir()
            assembly = cad_directory / "root.asm.1"
            assembly.write_bytes(b"assembly")
            run = RunRecord(
                run_id="run-1",
                bom_file=bom_file,
                cad_directory=cad_directory,
                workspace=workspace,
                status=RunStatus.NEEDS_REVIEW,
                input_fingerprint="sha256:input",
                plan_revision=1,
                created_at="2026-08-18T00:00:00+00:00",
                updated_at="2026-08-18T00:00:00+00:00",
            )
            store.add(run)
            store.replace_steps(
                run.run_id,
                (
                    StepResult(
                        "step-1",
                        "process-001",
                        StepStatus.PASSED,
                        (),
                        "sha256:committed-state-step-1",
                        None,
                    ),
                    StepResult(
                        "step-2",
                        "process-001",
                        StepStatus.QUESTIONED,
                        (),
                        "sha256:state-step-2",
                        None,
                    ),
                    StepResult(
                        "step-3",
                        "process-001",
                        StepStatus.QUESTIONED,
                        (),
                        "sha256:state-step-3",
                        revised_step_3_hash,
                    ),
                ),
            )
            artifacts = ArtifactStore(store)

            def write(kind: str, relative: str, value: dict) -> str:
                artifacts.write_json(
                    run_id=run.run_id,
                    run_workspace=workspace,
                    kind=kind,
                    relative_path=relative,
                    value=value,
                )
                return relative

            refs = (
                write(
                    "input-manifest",
                    "analysis/input-manifest.json",
                    {
                        "schema_version": "input-manifest/v1",
                        "bom": {
                            "name": bom_file.name,
                            "sha256": "sha256:"
                            + sha256(bom_file.read_bytes()).hexdigest(),
                        },
                        "cad": [
                            {
                                "relative_path": assembly.name,
                                "sha256": "sha256:"
                                + sha256(assembly.read_bytes()).hexdigest(),
                            }
                        ],
                    },
                ),
                write(
                    "normalized-bom",
                    "analysis/normalized-bom.json",
                    {
                        "schema_version": "normalized-bom/v1",
                        "sheet_name": "BOM",
                        "header_row": 1,
                        "columns": {"name": 1},
                        "rows": [
                            _bom_row(1, "ROOT", "总装"),
                            _bom_row(2, "A", "零件A"),
                            _bom_row(3, "B", "零件B"),
                            _bom_row(4, "C", "零件C"),
                        ],
                        "sheet_candidates": ["BOM"],
                    },
                ),
                write(
                    "formal-render-plan",
                    "plans/locked-render-plan-0001.json",
                    {
                        "schema_version": "formal-render-plan/v2",
                        "assembly_file": "root.asm.1",
                        "camera_basis": {},
                        "initial_completed_occurrences": [],
                        "scope_base_occurrences": {},
                        "steps": [
                            _formal_step("step-1", 2),
                            _formal_step("step-2", 3),
                            _formal_step("step-3", 4),
                        ],
                        "diagnostics": [],
                        "ready_steps": 2,
                        "questioned_steps": 0,
                        "checkpoint_interval": 20,
                        "scope_decisions": {},
                        "fingerprint": "sha256:plan",
                    },
                ),
                write(
                    "validation-result",
                    "results/validation-0001.json",
                    {
                        "steps": [
                            {
                                "step_id": "step-1",
                                "status": "QUESTIONED",
                                "image_path": "rendered/step-1.jpg",
                            },
                            {
                                "step_id": "step-2",
                                "status": "QUESTIONED",
                                "image_path": "rendered/step-2.jpg",
                                "complete_state_hash": "sha256:validated-state-step-2",
                                "depends_on": [],
                            },
                            {
                                "step_id": "step-3",
                                "status": "QUESTIONED",
                                "image_path": "rendered/step-3.jpg",
                                "complete_state_hash": "sha256:state-step-3",
                                "depends_on": [],
                            },
                        ]
                    },
                ),
                write(
                    "candidate-set",
                    "results/candidate-set-0001.json",
                    {"groups": []},
                ),
                write(
                    "publication-result",
                    "results/publication-0001.json",
                    {
                        "steps": [
                            {
                                "step_id": "step-1",
                                "status": "PASSED",
                                "image_path": "rendered/step-1.jpg",
                                "human_disposition": "accept_current_image",
                            },
                            {"step_id": "step-2", "status": "QUESTIONED"},
                            {
                                "step_id": "step-3",
                                "status": "QUESTIONED",
                                "image_path": "rendered/step-3-revision-0001.jpg",
                                "output_hash": revised_step_3_hash,
                            },
                        ]
                    },
                ),
                write(
                    "invalidation-set",
                    "revisions/invalidation-set-0002.json",
                    {"steps": ["step-2"]},
                ),
                write(
                    "step-revision",
                    "revisions/step-revision-0002.json",
                    {
                        "schema_version": "step-revision/v1",
                        "revision": 2,
                        "step_id": "step-2",
                        "kind": "presentation",
                        "changes": {
                            "candidate_id": SKIP_PUBLICATION_CANDIDATE_ID,
                        },
                    },
                ),
                write(
                    "human-review-decision",
                    "reviews/human-review-decision-0002.json",
                    {
                        "schema_version": "human-review-decision/v1",
                        "revision": 2,
                        "step_id": "step-2",
                        "decision": "skip_from_publication",
                        "reason": "当前图片无法满足要求，先出版其余步骤",
                        "publication_transform": "omitted",
                    },
                ),
            )
            output = publish_delivery(
                SkillContext(run, artifacts, store, {}),
                SkillInvocation(
                    "skill-invocation/v1",
                    run.run_id,
                    "publish-delivery",
                    refs,
                    {},
                ),
            )
            publication = output.artifacts[0].value
            by_step = {item["step_id"]: item for item in publication["steps"]}
            workbook = load_workbook(
                workspace / "delivery" / publication["workbook"],
                read_only=True,
                data_only=True,
            )
            published_text = "\n".join(
                str(cell.value)
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            )
            workbook.close()

        self.assertEqual(by_step["step-1"]["status"], "PASSED")
        self.assertEqual(
            by_step["step-1"]["complete_state_hash"],
            "sha256:committed-state-step-1",
        )
        self.assertEqual(by_step["step-2"]["status"], "PASSED")
        self.assertEqual(
            by_step["step-2"]["human_disposition"], "skip_from_publication"
        )
        self.assertFalse(by_step["step-2"]["included_in_sop"])
        self.assertNotIn("安装零件B", published_text)
        self.assertEqual(
            by_step["step-2"]["complete_state_hash"],
            "sha256:validated-state-step-2",
        )
        self.assertEqual(
            by_step["step-3"]["image_path"],
            "rendered/step-3-revision-0001.jpg",
        )
        self.assertEqual(by_step["step-3"]["output_hash"], revised_step_3_hash)


if __name__ == "__main__":
    unittest.main()
