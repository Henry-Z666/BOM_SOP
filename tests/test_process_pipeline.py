from __future__ import annotations

from dataclasses import asdict, replace
from pathlib import Path
import json
import tempfile
import unittest
import zipfile

from openpyxl import Workbook
from PIL import Image

from test_formal_render_planner import fixture
from sop_pipeline.agent.formal_render_planner import compile_formal_render_plan, lock_formal_render_plan
from sop_pipeline.agent.bom_normalizer import normalize_bom
from sop_pipeline.agent.model_inventory import ModelFile, matching_model_files_for_row
from sop_pipeline.agent.model_inventory import inventory_models
from sop_pipeline.agent.assembly_forest import build_assembly_forest
from sop_pipeline.agent.process_model import SourceEvidence, ProcessOperation, TorqueDecision, fingerprint, resolve_torque
from sop_pipeline.agent.process_sources import claims_in, read_process_card, read_source_package
from sop_pipeline.agent.process_bridge import OperationBinding, assert_process_sources, compile_process_jobs, enrich_sop_step
from sop_pipeline.agent.process_cli import compile_package, _locked_jobs
from sop_pipeline.agent.sop_publisher import SopStep, SopImage, SopPublisher
from sop_pipeline.agent.torque_validation import validate_torque_audit


def evidence(text='紧固 M6，力矩 4±0.4 N.m'):
    return SourceEvidence('工艺卡.xlsx', 'sha256:' + 'a'*64, '装配1', 'AJ8', text)


def geometry():
    bom, draft, mapping, graph = fixture()
    return lock_formal_render_plan(compile_formal_render_plan(bom, draft, mapping, graph),
        {'subassembly-scope-0006': '按BOM在本工位展开内部构造'})


def process_fixture():
    plan = geometry()
    operations, bindings = [], []
    for index, step in enumerate(plan.steps):
        item = evidence()
        claims = claims_in(item) if index == 0 else ()
        operation = ProcessOperation(f'op-{index}', f'操作{index}', 'pending', (), (item,), claims)
        decisions = (TorqueDecision(operation.operation_id, 'connection-1', claims[0].claim_id,
            step.moving_occurrences, 'test operator', 'fixture torque verified', (claims[0].claim_id,)),) if claims else ()
        binding = OperationBinding(operation.operation_id, (step.step_id,), 'test operator',
            'fixture product verified', decisions, bool(claims), '' if claims else 'positioning operation',
            'main_assembly' if step.stage_scope_occurrence == 'ROOT' else 'subassembly')
        operations.append(operation)
        bindings.append(binding)
    return plan, tuple(operations), tuple(bindings)


class ProcessPipelineTests(unittest.TestCase):
    def test_same_operation_only_labels_required_connection_steps(self):
        plan, operations, bindings = process_fixture()
        first, second = plan.steps[:2]
        second = replace(second, stage_scope_occurrence=first.stage_scope_occurrence)
        plan = replace(plan, steps=(first, second, *plan.steps[2:]))
        binding = replace(bindings[0], step_ids=(first.step_id, second.step_id),
            torque_target_occurrences=first.moving_occurrences)
        jobs = compile_process_jobs(plan, operations, (binding, *bindings[2:]))
        by_id = {task.step_id: task for task in jobs.tasks}
        self.assertEqual(by_id[first.step_id].payload['torque_annotations']['status'], 'confirmed')
        self.assertEqual(by_id[first.step_id].payload['torque_annotations']['labels'][0]['text'], '紧固力矩：4±0.4 N.m')
        self.assertEqual(by_id[second.step_id].payload['torque_annotations']['status'], 'not_applicable')
        self.assertEqual(by_id[second.step_id].payload['torque_annotations']['labels'], [])

    def test_parameters_retain_units_tolerance_and_source(self):
        item = evidence('初拧 10 N.M，复检 7 N.m；另一连接 40 kgf.cm')
        claims = claims_in(item)
        self.assertEqual([(c.value,c.unit) for c in claims], [('10','N.m'),('7','N.m'),('40','kgf.cm')])
        self.assertTrue(all(c.evidence == item for c in claims))
        self.assertEqual(claims_in(evidence())[0].text, '4±0.4 N.m')

    def test_conversion_typo_is_never_a_production_claim(self):
        self.assertEqual(claims_in(evidence('1kgf.cm=0.980665N.m≈0.098N.m')), ())
        self.assertEqual(claims_in(evidence('力矩 1 kgf.cm'))[0].value, '1')

    def test_conflicts_remain_pending_without_explicit_selection(self):
        item = evidence('卡箍 10 N.m；具体卡箍 6 N.m')
        op = ProcessOperation('op', '卡箍', 'pending', (), (item,), claims_in(item))
        self.assertIsNone(resolve_torque(op, None))
        decision = TorqueDecision('op', 'clamp', op.torque_claims[1].claim_id, ('20/2',),
            '工程师', '该连接按具体型号确认', tuple(c.claim_id for c in op.torque_claims))
        self.assertEqual(resolve_torque(op, decision).value, '6')
        with self.assertRaises(ValueError):
            resolve_torque(op, replace(decision, acknowledged_claim_ids=(decision.claim_id,)))

    def test_process_dependency_rebuilds_order_and_keeps_cad_dependencies(self):
        plan, operations, bindings = process_fixture()
        baseline = compile_process_jobs(plan, operations, bindings).fingerprint
        operations = (replace(operations[0], depends_on=(operations[2].operation_id,)), *operations[1:])
        jobs = compile_process_jobs(plan, operations, bindings)
        self.assertEqual(jobs.tasks[-1].step_id, plan.steps[0].step_id)
        original = next(step for step in plan.steps if step.stage_scope_occurrence == 'ROOT')
        compiled = next(task for task in jobs.tasks if task.step_id == original.step_id)
        self.assertTrue(set(original.depends_on).issubset(compiled.depends_on))
        self.assertNotEqual(jobs.fingerprint, baseline)
        # Independent subassembly scope stays isolated even when global chronology changes.
        self.assertEqual(jobs.tasks[-1].complete_state_hash, plan.steps[0].complete_state_hash)

    def test_process_cycle_cannot_be_silently_reordered(self):
        plan, operations, bindings = process_fixture()
        operations = (replace(operations[0], depends_on=(operations[1].operation_id,)),
            replace(operations[1], depends_on=(operations[0].operation_id,)), operations[2])
        with self.assertRaises(ValueError):
            compile_process_jobs(plan, operations, bindings)

    def test_every_step_requires_confirmed_scope_and_single_owner(self):
        plan, operations, bindings = process_fixture()
        for invalid in (bindings[:-1], (replace(bindings[0], scope='pending'), *bindings[1:]),
                (replace(bindings[0], scope='main_assembly'), *bindings[1:])):
            with self.assertRaises(ValueError):
                compile_process_jobs(plan, operations, invalid)

    def test_unknown_or_duplicate_torque_position_is_rejected(self):
        plan, operations, bindings = process_fixture()
        decision = bindings[0].torque_decisions[0]
        for decisions in ((replace(decision, occurrence_ids=('999/1',)),), (decision,decision)):
            with self.assertRaises(ValueError):
                compile_process_jobs(plan, operations, (replace(bindings[0], torque_decisions=decisions), *bindings[1:]))

    def test_missing_torque_cannot_be_bypassed_by_formal_publication(self):
        plan, operations, bindings = process_fixture()
        bindings = (replace(bindings[0], torque_decisions=()), *bindings[1:])
        payload = compile_process_jobs(plan, operations, bindings).tasks[0].payload
        self.assertEqual(payload['publication_parameter_status'], 'pending')
        step = SopStep('step','p','p','t',SopImage('image',Path('unused.jpg')),(),'','','')
        step = enrich_sop_step(step, payload)
        self.assertIn('待确认', step.control_points)
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder)/'delivery'
            with self.assertRaises(ValueError):
                SopPublisher().publish((step,),output)
            self.assertFalse(output.exists())

    def test_publication_preserves_interleaved_process_dependencies(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder); image=root/'test.jpg';Image.new('RGB',(320,240),'white').save(image)
            steps=tuple(SopStep(f's-{index}',process,process,f'action-{index}',SopImage(f'i-{index}',image),
                (),f'action-{index}','','',preserve_process_order=True) for index,process in enumerate(('A','B','A')))
            output=SopPublisher().publish(steps,root/'delivery')
            workbook=load_workbook(output)
            self.assertEqual(workbook.sheetnames,['A','B','A-续页2'])
            self.assertEqual([len(sheet._images) for sheet in workbook],[1,1,1])
            workbook.close()

    def test_native_audit_requires_exact_coverage_and_manual_visual_review(self):
        plan, operations, bindings = process_fixture()
        payload = compile_process_jobs(plan, operations, bindings).tasks[0].payload
        contract = payload['torque_annotations']
        selected = contract['labels'][0]
        with tempfile.TemporaryDirectory() as folder:
            image = Path(folder)/'step.jpg'
            self.assertEqual(validate_torque_audit(image,payload), ('TORQUE_AUDIT_MISSING',))
            audit = {'schema_version':'torque-label-audit/v1','renderer':'creo_display_list/v1',
                'contract_hash':contract['contract_hash'],'overlap_check':'passed','labels':[
                    {'occurrence_id':selected['occurrence_ids'][0],'connection_id':selected['connection_id'],
                     'text':selected['text'].replace('±','+/-'),'anchor_source':'native_solid_surface',
                     'target_root':[0,0,0],'text_box_camera':[10,10,40,10]}]}
            path=Path(str(image)+'.torque.json')
            path.write_text(json.dumps(audit),encoding='utf-8')
            self.assertEqual(validate_torque_audit(image,payload),('TORQUE_VISUAL_REVIEW_REQUIRED',))
            audit['labels'][0]['occurrence_id']='other'
            path.write_text(json.dumps(audit),encoding='utf-8')
            self.assertEqual(validate_torque_audit(image,payload),('TORQUE_COVERAGE_MISMATCH',))

    def test_qsl_adapter_reads_actual_text_box_relationships(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'工艺卡.xlsx'
            book=Workbook(); sheet=book.active; sheet.title='连接安装'
            sheet['AJ8']='卡箍安装'; sheet['AJ21']='123456789012'; book.save(path); book.close()
            with zipfile.ZipFile(path) as original:
                files={name:original.read(name) for name in original.namelist()}
            files['xl/worksheets/sheet1.xml']=files['xl/worksheets/sheet1.xml'].replace(b'</worksheet>',
                b'<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/></worksheet>')
            files['xl/worksheets/_rels/sheet1.xml.rels']=b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="../drawings/drawing1.xml"/></Relationships>'
            files['xl/drawings/drawing1.xml']=('<wsDr xmlns="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
                'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><twoCellAnchor><a:t>连接力矩 6 N.M</a:t></twoCellAnchor></wsDr>').encode()
            with zipfile.ZipFile(path,'w') as target:
                for name,data in files.items(): target.writestr(name,data)
            operation=read_process_card(path)[0]
            self.assertEqual(operation.materials,('123456789012',))
            self.assertEqual(operation.torque_claims[0].value,'6')
            self.assertIn('#anchor-',operation.torque_claims[0].evidence.location)
            package=read_source_package((path,))
            assert_process_sources((operation,),tuple(package['sources']))
            forged=replace(operation,instructions=(replace(operation.instructions[0],text='伪造'),))
            with self.assertRaises(ValueError): assert_process_sources((forged,),tuple(package['sources']))

    def test_new_bom_columns_and_multiline_cad_names_are_preserved(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'BOM.xlsx'; book=Workbook(); sheet=book.active
            sheet.append(['层次','图号','名称','型号','模型名称','对应工艺路线工序','发料\n工位'])
            sheet.append(['1.1','drawing','紧固件','generic','bolt-a\nbolt-b','框架安装','框架安装'])
            book.save(path); book.close()
            row=normalize_bom(path).rows[0]
            self.assertEqual(row.station,'框架安装')
            self.assertEqual(row.process_route,'框架安装')
            files=(ModelFile('bolt-a.prt.1','bolt-a','prt',1),ModelFile('bolt-b.prt.1','bolt-b','prt',1))
            self.assertEqual(matching_model_files_for_row(row,files),files)

    def test_explicit_product_root_preserves_source_level_and_owns_all_rows(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder); path=root/'BOM.xlsx'; book=Workbook(); sheet=book.active
            sheet.append(['层次','图号','名称','模型名称'])
            sheet.append(['','PRODUCT','产品','product'])
            sheet.append(['1','FRAME','框架','frame'])
            sheet.append(['1.1','BOLT','螺钉','bolt'])
            book.save(path);book.close()
            for name in ('product.asm.3','frame.asm.8','bolt.prt.1'):
                (root/name).write_bytes(b'fixture; not native CAD')
            bom=normalize_bom(path,root_bom_row=2)
            self.assertEqual([row.source_level for row in bom.rows],['','1','1.1'])
            self.assertEqual([row.level for row in bom.rows],['0','0.1','0.1.1'])
            inventory=inventory_models(root,bom)
            self.assertEqual(inventory.final_assembly,'product.asm.3')
            forest=build_assembly_forest(bom,inventory)
            self.assertEqual(len(forest.roots),1)
            self.assertEqual(forest.roots[0].source_bom_rows,(2,3,4))

    def test_cli_compilation_locks_source_package_and_decisions(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder); source=root/'card.xlsx'; workbook=Workbook()
            workbook.active.title='紧固安装';workbook.active['AJ8']='力矩 4 N.m'; workbook.save(source);workbook.close()
            package=read_source_package((source,))
            parsed=read_process_card(source)[0]
            plan,operations,bindings=process_fixture()
            adapted=[];confirmed=[]
            for operation,binding in zip(operations,bindings):
                updated=replace(operation,instructions=parsed.instructions,
                    torque_claims=parsed.torque_claims if binding.torque_required else ())
                adapted.append(asdict(updated))
                decisions=tuple(replace(decision,claim_id=parsed.torque_claims[0].claim_id,
                    acknowledged_claim_ids=(parsed.torque_claims[0].claim_id,)) for decision in binding.torque_decisions)
                confirmed.append(asdict(replace(binding,torque_decisions=decisions)))
            package['operations']=adapted
            package_path=root/'package.json'; package_path.write_text(json.dumps(package),encoding='utf-8')
            geometry_path=root/'geometry.json';geometry_path.write_text(json.dumps(asdict(plan)),encoding='utf-8')
            decisions_path=root/'decisions.json';decisions_path.write_text(json.dumps({
                'source_package_fingerprint':package['source_package_fingerprint'],
                'geometry_fingerprint':plan.fingerprint,'bindings':confirmed}),encoding='utf-8')
            jobs_path=root/'jobs.json'
            compiled=compile_package(package_path,geometry_path,decisions_path,jobs_path)
            self.assertEqual(len(compiled['tasks']),3)
            self.assertEqual(len(_locked_jobs(jobs_path)[1].tasks),3)
            source.write_bytes(b'changed')
            with self.assertRaises(ValueError): _locked_jobs(jobs_path)


if __name__ == '__main__':
    unittest.main()
