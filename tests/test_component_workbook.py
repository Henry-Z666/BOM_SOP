from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook
from PIL import Image

from sop_pipeline.component_workbook import build_component_workbook, verify_workbook

ASSETS = Path(__file__).resolve().parents[1]/'assets/publication'


class ComponentWorkbookTests(unittest.TestCase):
    def fixture(self, output):
        image = output/'source.jpg'
        Image.new('RGB', (100, 100), 'red').save(image)
        digest = sha256(image.read_bytes()).hexdigest()
        images = [{'task_id': f'i{n}', 'image_path': str(image)} for n in range(1, 4)]
        material = {'row': 4, 'material_code': '000123', 'name': '完整来料',
                    'model': '型号', 'quantity': 2, 'unit': '件', 'assembly_boundary': True}
        pages = [{'number': n, 'package_number': n, 'part': 1, 'part_count': 1,
            'component_title': f'部件{n}', 'chapter': '安装', 'materials': [deepcopy(material)],
            'instructions': ['按来源安装，紧固力矩 10 N·m'],
            'images': [f'i{i}' for i in range(1, n+1)],
            'captions': {f'i{i}': f'连接{i}' for i in range(1, n+1)},
            'actions': ['安装整体', '完成连接'], 'tools': [{'name': '扭矩扳手', 'quantity': 1}]}
            for n in range(1, 4)]
        data = {'publication_mode': 'component_cards/v1', 'publication_file': 'SOP.xlsx',
                'publication_title': '装配工艺卡', 'publication_code': 'TEST',
                'publication_date': '2026-10-30', 'images': images, 'pages': pages}
        gates = [{'task_id': f'i{n}', 'passed': True, 'image_sha256': digest} for n in range(1, 4)]
        (output/'出版内容.json').write_text(json.dumps(data, ensure_ascii=False), encoding='utf8')
        (output/'出图检查.json').write_text(json.dumps(gates), encoding='utf8')
        return data, image

    def test_one_two_three_image_pages_preserve_bytes_tables_and_printing(self):
        with TemporaryDirectory() as directory:
            output = Path(directory)
            data, image = self.fixture(output)
            target = build_component_workbook(output, ASSETS)
            report = verify_workbook(target, data)
            self.assertEqual(report['pages'], 3)
            self.assertTrue(report['embedded_image_bytes_unchanged'])
            workbook = load_workbook(target)
            try:
                for index, sheet in enumerate(workbook.worksheets, 1):
                    self.assertEqual(len(sheet._images), index)
                    self.assertEqual(sheet['AJ21'].value, '000123')
                    self.assertEqual(sheet['AR21'].value, 2)
                    self.assertEqual(sheet['AM21'].fill.fgColor.rgb[-6:], 'FFFF99')
                    self.assertEqual(sheet.page_setup.orientation, 'landscape')
                    self.assertEqual(str(sheet.page_setup.paperSize), sheet.PAPERSIZE_A4)
                    self.assertIn('B31:Q35', set(map(str, sheet.merged_cells.ranges)))
                    for drawing in sheet._images:
                        self.assertLess(drawing.anchor.to.col, 34)
                        self.assertLess(drawing.anchor.to.row, 31)
            finally:
                workbook.close()

    def test_changed_or_missing_source_image_does_not_replace_previous_workbook(self):
        with TemporaryDirectory() as directory:
            output = Path(directory)
            data, image = self.fixture(output)
            target = build_component_workbook(output, ASSETS)
            previous = target.read_bytes()
            Image.new('RGB', (100, 100), 'blue').save(image)
            with self.assertRaisesRegex(ValueError, 'bytes changed'):
                build_component_workbook(output, ASSETS)
            self.assertEqual(target.read_bytes(), previous)
            image.unlink()
            with self.assertRaisesRegex(FileNotFoundError, 'Copy the local image package'):
                build_component_workbook(output, ASSETS)

    def test_update_selected_page_preserves_other_pages(self):
        with TemporaryDirectory() as directory:
            output = Path(directory)
            data, _ = self.fixture(output)
            target = build_component_workbook(output, ASSETS)
            data['pages'][1]['instructions'][0] = '仅更新第二页'
            (output/'出版内容.json').write_text(json.dumps(data), encoding='utf8')
            build_component_workbook(output, ASSETS, update_pages=[2])
            workbook = load_workbook(target)
            try:
                self.assertEqual(workbook.worksheets[1]['AJ8'].value, '仅更新第二页')
                self.assertEqual(workbook.worksheets[0]['AJ8'].value, '按来源安装，紧固力矩 10 N·m')
                self.assertEqual([len(s._images) for s in workbook], [1, 2, 3])
            finally:
                workbook.close()

    def test_over_capacity_page_is_rejected(self):
        with TemporaryDirectory() as directory:
            output = Path(directory)
            data, _ = self.fixture(output)
            data['pages'][0]['materials'] *= 10
            (output/'出版内容.json').write_text(json.dumps(data), encoding='utf8')
            with self.assertRaisesRegex(ValueError, 'capacity'):
                build_component_workbook(output, ASSETS)
