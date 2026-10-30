"""Portable component-card XLSX publication using openpyxl, without AI services."""
from __future__ import annotations

from collections import Counter
from hashlib import sha256
from pathlib import Path
import json
import math
import re

from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU


def _read(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def _material_text(material):
    unit = material.get('unit', '')
    suffix = f'（单位：{unit}）' if unit and unit not in {'个', '件', '台', '套', '只'} else ''
    return material['name'] + suffix + '\n' + (material.get('model') or material.get('drawing_no', ''))


def _lines(text, capacity=28):
    return sum(max(1, math.ceil(sum(2 if ord(c) > 255 else 1 for c in line) / capacity))
               for line in text.split('\n'))


def validate_assets(data, gates, image_ids=None):
    """Check every required native image's bytes before opening the output file."""
    images = {image['task_id']: image for image in data['images']}
    if len(images) != len(data['images']):
        raise ValueError('Duplicate image task ID')
    gate_map = {gate['task_id']: gate for gate in gates}
    if len(gate_map) != len(gates) or set(gate_map) != set(images):
        raise ValueError('Image check records do not match image tasks')
    missing = []
    for identity in (images if image_ids is None else image_ids):
        if identity not in images or identity not in gate_map:
            raise ValueError('Unknown image task: ' + identity)
        gate, image = gate_map[identity], images[identity]
        if not gate.get('passed'):
            raise ValueError('Image did not pass native checks: ' + identity)
        path = Path(image['image_path'])
        if not path.is_file():
            missing.append(str(path))
        elif sha256(path.read_bytes()).hexdigest() != gate['image_sha256'].removeprefix('sha256:'):
            raise ValueError('Native image bytes changed: ' + identity)
    if missing:
        raise FileNotFoundError(f'Missing {len(missing)} native image(s). Copy the local image package first. '
                                f'First missing file: {missing[0]}')


def _add_image(sheet, path, x, y, width, height, widths, heights, row=7, column=1):
    def marker(offset_x, offset_y):
        c, r = column, row
        while c < len(widths)-1 and offset_x >= widths[c]:
            offset_x -= widths[c]
            c += 1
        while r < len(heights)-1 and offset_y >= heights[r]:
            offset_y -= heights[r]
            r += 1
        return AnchorMarker(col=c, row=r, colOff=pixels_to_EMU(offset_x), rowOff=pixels_to_EMU(offset_y))
    image = Image(path)
    image.anchor = TwoCellAnchor(editAs='oneCell', _from=marker(x, y), to=marker(x+width, y+height))
    sheet.add_image(image)


def _text(sheet, address, value, *, size=10, bold=False, color='243442', fill=None, center=False):
    if ':' in address:
        sheet.merge_cells(address)
    cell = sheet[address.split(':')[0]]
    cell.value = value
    cell.font = Font(name='Microsoft YaHei', size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal='center' if center else 'left',
                               vertical='center', wrap_text=True)
    if fill:
        for cells in sheet[address] if ':' in address else ((cell,),):
            for item in cells:
                item.fill = PatternFill('solid', fgColor=fill)


def _populate(sheet, page, data, images, meta, heights_override, total_pages):
    sheet._images = []
    sheet.sheet_view.showGridLines = False
    widths = [math.floor(meta.get('defaultColWidth', 3.5) * 7 + 5)] * 44
    for dimension in meta['cols'].values():
        for index in range(dimension['min'], min(44, dimension['max']) + 1):
            widths[index - 1] = math.floor(dimension['width'] * 7 + 5)
    # Remove the template's grouped column spans before defining single columns.
    # Otherwise overlapping <col min/max> records can change image/table geometry.
    sheet.column_dimensions.clear()
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = (width - 5) / 7
    heights = []
    for index in range(35):
        points = meta['rows'].get(str(index + 1)) or meta['defaultRowHeight']
        if 7 <= index < 7 + len(page['instructions']):
            points = max(points, _lines(page['instructions'][index - 7], 52) * 12 + 6)
        if index in {7, 21}:
            points = max(points, 28)
        if index == 4:
            points = max(points, 36)
        if index >= 30:
            points = max(points, 23)
        if 20 <= index < 20 + len(page['materials']):
            points = max(points, heights_override.get(str(index + 1), 0), 27,
                         _lines(_material_text(page['materials'][index - 20])) * 12 + 6)
        sheet.row_dimensions[index + 1].height = points
        heights.append(points * 4 / 3)
    fields = {'J1': data.get('publication_title', '装配工艺卡'),
        'B4': f"文件编号：{data.get('publication_code', 'SOP')} / {page['number']}–{total_pages}",
        'M4': data.get('product_model', 'CDU900-1250L-S-D7L'), 'V4': '',
        'AF1': data.get('publication_date', ''), 'AN1': 'C1', 'AN3': '01',
        'AN4': f"ZZ.{page['package_number']}-{page['part']}", 'AN5': page['component_title'],
        'AN6': '', 'B7': page['component_title'] +
        (f" · {page['section']}" if page.get('section') else '') +
        (f"（{page['part']}/{page['part_count']}）" if page['part_count'] > 1 else '')}
    for address, value in fields.items():
        _text(sheet, address, value, size=14 if address == 'J1' else 9, bold=address == 'J1')
    for address in ('AF2', 'AF3', 'AF4', 'AF5', 'AF6'):
        sheet[address] = ''
    for number in range(8, 19):
        sheet[f'AI{number}'] = number - 7
        _text(sheet, f'AJ{number}', page['instructions'][number - 8]
              if number - 8 < len(page['instructions']) else None, size=9, color='16202A')
    for number in range(21, 30):
        for column in ('AI', 'AJ', 'AM', 'AR'):
            sheet[f'{column}{number}'] = None
    for number, material in enumerate(page['materials'], 21):
        sheet[f'AI{number}'] = number - 20
        _text(sheet, f'AJ{number}', material['material_code'], size=9)
        sheet[f'AJ{number}'].number_format = '0'
        _text(sheet, f'AM{number}', _material_text(material), size=9)
        sheet[f'AR{number}'] = material['quantity']
        if material.get('assembly_boundary'):
            for cells in sheet[f'AM{number}:AQ{number}']:
                for cell in cells:
                    cell.fill = PatternFill('solid', fgColor='FFFF99')
    for number in range(32, 36):
        for column in ('AI', 'AJ', 'AM', 'AR'):
            sheet[f'{column}{number}'] = None
    for number, tool in enumerate(page.get('tools', []), 32):
        sheet[f'AI{number}'] = number - 31
        _text(sheet, f'AJ{number}', tool['name'], size=9)
        _text(sheet, f'AM{number}', tool.get('spec', ''), size=9)
        sheet[f'AR{number}'] = tool.get('quantity', '')
    if data.get('branding'):
        header_width, header_height = sum(widths[1:9]), sum(heights[:3])
        width = min(header_width - 24, header_height * 1043 / 292 - 12)
        height = width * 292 / 1043
        _add_image(sheet, data['branding']['image_path'], (header_width - width) / 2,
            (header_height - height) / 2, width, height, widths, heights, row=0)
    zone_width, figure_height = sum(widths[1:34]), sum(heights[7:30])
    def insert(identity, x, y, width, height):
        _add_image(sheet, images[identity]['image_path'], x, y, width, height, widths, heights)
    def caption(address, identity):
        _text(sheet, address, page['captions'][identity], size=10, bold=True, center=True)
    count = len(page['images'])
    if count == 3:
        left = sum(widths[1:22])
        right = zone_width - left
        split = sum(heights[7:19])
        caption('B8:V8', page['images'][0])
        caption('W8:AH8', page['images'][1])
        caption('W20:AH20', page['images'][2])
        size = min(left - 18, figure_height - heights[7] - 16)
        insert(page['images'][0], (left-size)/2,
               heights[7]+(figure_height-heights[7]-size)/2, size, size)
        for index in (1, 2):
            top, band, head = (0, split, heights[7]) if index == 1 else (split, figure_height-split, heights[19])
            size = min(right-18, band-head-10)
            insert(page['images'][index], left+(right-size)/2, top+head+(band-head-size)/2, size, size)
    elif count:
        slot = zone_width / count
        size = min(slot-18, figure_height-heights[7]-16)
        for index, identity in enumerate(page['images']):
            caption('B8:AH8' if count == 1 else 'B8:Q8' if index == 0 else 'R8:AH8', identity)
            insert(identity, index*slot+(slot-size)/2,
                   heights[7]+(figure_height-heights[7]-size)/2, size, size)
    actions = page.get('actions', [])
    ranges = (['B31:L35', 'M31:W35', 'X31:AH35'] if len(actions) == 3 else
              ['B31:Q35', 'R31:AH35'] if len(actions) == 2 else ['B31:AH35'])
    for index, action in enumerate(actions):
        _text(sheet, ranges[index], f'{index+1}　{action}', size=11, bold=True,
              color='382B13', fill='FFF0B3')
    if not count and not actions:
        _text(sheet, 'B10:AH27', '\n\n'.join(page['instructions']), size=13, color='16202A')
    sheet.print_area = 'B1:AR35'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.print_options.gridLines = False
    sheet.print_options.headings = False
    sheet.page_margins.left = sheet.page_margins.right = .15
    sheet.page_margins.top = sheet.page_margins.bottom = .2
    sheet.page_margins.header = sheet.page_margins.footer = 0


def verify_workbook(path: Path, data: dict, pages=None) -> dict:
    pages = data['pages'] if pages is None else pages
    workbook = load_workbook(path)
    try:
        if len(workbook.worksheets) != len(data['pages']):
            raise ValueError('Worksheet count differs from publication pages')
        images = {entry['task_id']: entry for entry in data['images']}
        for page in pages:
            sheet = workbook.worksheets[page['number'] - 1]
            files = [Path(images[identity]['image_path']) for identity in page['images']]
            if data.get('branding'):
                files.append(Path(data['branding']['image_path']))
            expected = Counter(sha256(file.read_bytes()).hexdigest() for file in files)
            actual = Counter(sha256(image._data()).hexdigest() for image in sheet._images)
            if actual != expected:
                raise ValueError(f"Embedded image bytes differ on page {page['number']}")
            if ('B1:AR35' not in str(sheet.print_area).replace('$', '') or
                sheet.page_setup.fitToHeight != 1 or sheet.page_setup.fitToWidth != 1 or
                sheet.page_setup.orientation != 'landscape' or
                str(sheet.page_setup.paperSize) != sheet.PAPERSIZE_A4):
                raise ValueError('Print layout differs from A4 page contract')
            for number, material in enumerate(page['materials'], 21):
                if (str(sheet[f'AJ{number}'].value) != material['material_code'] or
                    sheet[f'AR{number}'].value != material['quantity'] or
                    sheet[f'AM{number}'].value != _material_text(material)):
                    raise ValueError(f"Material table differs on page {page['number']}")
            for number, instruction in enumerate(page['instructions'], 8):
                if sheet[f'AJ{number}'].value != instruction:
                    raise ValueError(f"Instruction text differs on page {page['number']}")
            if any(cell.data_type == 'e' for cells in sheet for cell in cells):
                raise ValueError('Workbook contains an Excel error cell')
        return {'pages': len(data['pages']), 'verified_pages': len(pages),
                'operations': data.get('operation_count'),
                'material_occurrences': sum(len(m.get('occurrence_ids', []))
                    for page in data['pages'] for m in page['materials']),
                'native_images': len(data['images']), 'embedded_image_bytes_unchanged': True,
                'workbook_sha256': sha256(path.read_bytes()).hexdigest(),
                'printing': 'A4 landscape, one worksheet per printed page',
                'publisher': 'openpyxl', 'requires_openai': False}
    finally:
        workbook.close()


def build_component_workbook(output: Path, assets: Path, update_pages=None) -> Path:
    data = _read(output / '出版内容.json')
    if data.get('publication_mode') != 'component_cards/v1':
        raise ValueError('This publisher requires component_cards/v1')
    pages = data['pages']
    if [page['number'] for page in pages] != list(range(1, len(pages)+1)):
        raise ValueError('Page numbers must be consecutive and start at 1')
    if update_pages and (set(update_pages) - {p['number'] for p in pages}):
        raise ValueError('Unknown update page number')
    selected = [page for page in pages if not update_pages or page['number'] in update_pages]
    for page in selected:
        if (len(page['materials']) > 9 or len(page['instructions']) > 11 or
            len(page.get('tools', [])) > 4 or len(page['images']) > 3 or
            len(page.get('actions', [])) > 3 or page.get('reference_images')):
            raise ValueError('Page exceeds the component-card template capacity')
    validate_assets(data, _read(output / '出图检查.json'))
    filename = data['publication_file']
    target = (output / filename).resolve()
    if not target.is_relative_to(output.resolve()) or target.suffix.lower() != '.xlsx':
        raise ValueError('Publication file must be an XLSX inside the output directory')
    workbook = load_workbook(target if update_pages else assets / 'component-sop-template.xlsx')
    meta = _read(assets / 'template-meta.json')
    overrides_file = output / '表格行高修正.json'
    overrides = _read(overrides_file) if overrides_file.exists() else {}
    try:
        template = workbook.worksheets[0]
        if not update_pages:
            for page in pages[1:]:
                workbook.copy_worksheet(template)
        for page in selected:
            sheet = workbook.worksheets[page['number'] - 1]
            for merge in list(sheet.merged_cells.ranges):
                sheet.unmerge_cells(str(merge))
            for merge in meta['merges']:
                sheet.merge_cells(merge)
            sheet.merge_cells('AM21:AQ21')
            sheet.title = re.sub(r'[\\/*?:\[\]]', '_', f"{page['number']:03} {page['component_title']}")[:31]
            _populate(sheet, page, data, {image['task_id']: image for image in data['images']},
                      meta, overrides.get(str(page['number']), {}), len(pages))
        temporary = target.with_suffix('.building.xlsx')
        workbook.save(temporary)
    finally:
        workbook.close()
    try:
        report = verify_workbook(temporary, data)
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)
    (output / 'SOP交付检查.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf8')
    return target
